"""Importação das planilhas da equipe, em duas etapas: (1) leitura e pré-visualização com sinalização de
inconsistências; (2) confirmação pelo administrador, gravando tudo em uma única transação.

Princípios: nada é inventado — dado ausente ou ambíguo fica vazio e sinalizado para revisão; células vazias,
legendas, cabeçalhos e fórmulas não viram registros; os arquivos originais ficam guardados; reimportar não duplica
(cada registro leva a chave de origem em `source`)."""
import hashlib
import io
import json
import re
from datetime import date, datetime

import openpyxl

from ..core import (STAFF, Ctx, F, audit, bad, conflict, forbidden, fmt_date, from_cents, get_setting, not_found, route, today_local)
from ..db import new_id, now_iso
from .scope import norm

KIND_LABEL = {"gantt": "Matriz Gantt", "planejamento": "Planejamento de gestão", "raci": "Matriz RACI", "riscos": "Matriz de riscos",
              "eap": "Estrutura Analítica do Projeto", "orcamento": "Orçamento", "stakeholders": "Stakeholders", "escopo": "Escopo do projeto",
              "evidencias": "Evidências de monitoramento"}
SECTION_LABEL = {"people": "Integrantes", "wbs": "EAP", "activities": "Atividades", "risks": "Riscos", "expenses": "Itens do orçamento",
                 "stakeholders": "Stakeholders", "scope": "Escopo", "evidence": "Evidências"}
AREA_ALIASES = {"engenharia": "Engenharia", "engenheiro": "Engenharia", "gestao": "Gestão do projeto", "gestao do projeto": "Gestão do projeto",
                "empreendimento": "Empreendimento", "empreedimento": "Empreendimento", "empreendedorismo": "Empreendimento", "projeto social": "Projeto social"}
PROB_MAP = {"raro": 1, "improvavel": 2, "possivel": 3, "provavel": 4, "muito provavel": 5}
IMPACT_MAP = {"muito baixo": 1, "baixo": 2, "medio": 3, "alto": 4, "muito alto": 5}
MONTH_NUM = {m: i for i, m in enumerate(["janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"], 1)}
STOP = {"de", "do", "da", "dos", "das", "e", "o", "a", "os", "as"}
MILESTONE_COLORS = {"CC4125", "3C78D8", "674EA7", "358F0E"}


def area_of(text):
    return AREA_ALIASES.get(norm(text).rstrip(" :-").strip())


def tokens(s):
    return frozenset(t for t in norm(s).replace("gannt", "gantt").split() if t not in STOP)


def fill_of(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return None
    c = f.fgColor
    return c.rgb[-6:].upper() if c is not None and c.type == "rgb" and isinstance(c.rgb, str) else None


def txt(v):
    if v is None:
        return ""
    return str(v).strip()


def flag(level, text):
    return {"level": level, "text": text}


class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.rows = ws.max_row
        self.cols = ws.max_column

    def v(self, r, c):
        return self.ws.cell(r, c).value

    def s(self, r, c):
        return txt(self.ws.cell(r, c).value)


def detect(wb):
    ws = wb.worksheets[0]
    texts = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        for c in row:
            if isinstance(c.value, str) and c.value.strip():
                texts.append(norm(c.value))
    joined = " | ".join(texts)
    short = [t for t in texts if len(t) <= 60]

    first = texts[0] if texts else ""

    def title(*starts):
        # o título é a primeira célula preenchida; só na falta dele procura nas demais células curtas
        if any(first.startswith(st) for st in starts):
            return True
        return not known_first and any(t.startswith(st) for t in short for st in starts)
    known_first = any(first.startswith(k) for k in ("evidencias de monitoramento", "escopo do projeto", "planejamento de gestao", "cronograma",
                                                    "matriz e plano", "matriz raci", "orcamento", "registro de", "1 "))
    # a ordem importa: títulos mais específicos primeiro (textos longos podem citar outras planilhas)
    if title("evidencias de monitoramento"):
        return "evidencias"
    if title("escopo do projeto"):
        return "escopo"
    if title("planejamento de gestao"):
        return "planejamento"
    if title("cronograma") and "dependencia" in short:
        return "gantt"
    if title("matriz e plano de gerenciamento de riscos") or (title("risco") and "probabilidade" in short and "impacto" in short):
        return "riscos"
    if title("matriz raci"):
        return "raci"
    if title("orcamento") and "preco unitario" in short:
        return "orcamento"
    if title("registro de steackholders", "registro de stakeholders", "possiveis patrocinadores"):
        return "stakeholders"
    if sum(1 for t in texts if re.match(r"^1 \d", t)) >= 3:
        return "eap"
    return None


# ----------------------------------------------------------------------------- leitores por planilha
_RANGE_SEP = re.compile(r"\s[-–—]\s|[–—]")


def parse_range(value, year):
    """'05/06 - 30/06' -> (início, término, problemas). Aceita data única (vira prazo final)."""
    problems = []
    if isinstance(value, (datetime, date)):
        dt = value.date() if isinstance(value, datetime) else value
        return None, dt.isoformat(), problems
    s = txt(value).replace("\\", "/")
    if not s:
        return None, None, problems
    found = re.findall(r"(\d{1,2})\s*/\s*(\d{1,2})", s)
    out = []
    for dd, mm in found:
        try:
            out.append(date(year, int(mm), int(dd)).isoformat())
        except ValueError:
            out.append(None)
            problems.append(f"data inexistente “{dd}/{mm}”")
    if not found:
        problems.append(f"período não reconhecido “{s}”")
        return None, None, problems
    if len(out) >= 2:
        return out[0], out[1], problems
    # uma só data: se há separador depois dela é apenas o início (ex.: "15/08 – ")
    if re.search(r"\d\s*[-–—]\s*$", s):
        return out[0], None, problems
    return None, out[0], problems


def read_gantt(wb, year, today):
    ws = wb.worksheets[0]
    sh = Sheet(ws)
    items, notes = [], []
    m = re.search(r"(20\d{2})", " ".join(sh.s(r, 2) for r in range(1, 4)))
    if m:
        year = int(m.group(1))
    area = None
    future_real = 0
    for r in range(1, sh.rows + 1):
        e, fcol = sh.v(r, 5), sh.s(r, 6)
        if isinstance(e, str) and norm(fcol) == "dependencia":
            area = area_of(e)
            if not area:
                notes.append(flag("revisar", f"Linha {r}: área “{e.strip()}” não reconhecida."))
            continue
        if not isinstance(e, str) or e.startswith("="):
            continue
        mt = re.match(r"^\s*(\d+)\.\s*(.+?)\s*$", e.strip(), re.S)
        if not mt or not area:
            continue
        seq, title = int(mt.group(1)), re.sub(r"\s+", " ", mt.group(2)).strip()
        flags = []
        dep_raw = sh.v(r, 6)
        deps = [int(x) for x in re.findall(r"\d+", str(int(dep_raw)) if isinstance(dep_raw, float) else txt(dep_raw))] if dep_raw is not None else []
        ps, pe, pr = parse_range(sh.v(r, 7), year)
        as_, ae, ar = parse_range(sh.v(r, 8), year)
        for p in pr:
            flags.append(flag("revisar", f"Tempo estimado: {p}."))
        for p in ar:
            flags.append(flag("revisar", f"Tempo real: {p} — data real não importada."))
        if ps and pe and ps > pe:
            flags.append(flag("revisar", "Início planejado posterior ao término; início não importado."))
            ps = None
        if not pe and ps:
            pe, ps = ps, None
        if not pe:
            flags.append(flag("revisar", "Sem prazo planejado na planilha."))
        elif not ps:
            flags.append(flag("info", "A planilha traz apenas o prazo final (sem data de início)."))
        if as_ and as_ > today:
            as_, ae = None, None
            future_real += 1
        elif ae and ae > today:
            ae = None
            future_real += 1
        if as_ and ae and as_ > ae:
            flags.append(flag("revisar", "Início real posterior à conclusão real; conclusão não importada."))
            ae = None
        if sh.v(r, 8) is None and pe and pe <= today:
            flags.append(flag("revisar", "Sem “tempo real” na planilha: a situação não pôde ser determinada; confirme se a entrega foi concluída."))
        status = "concluida" if ae else "em_andamento" if as_ else "nao_iniciada"
        milestone = any(fill_of(ws.cell(rr, c)) in MILESTONE_COLORS for rr in range(r, min(r + 4, sh.rows + 1)) for c in range(9, 37))
        items.append({"key": f"gantt:{norm(area)}:{seq}", "origin": "gantt", "area": area, "seq": seq, "title": title, "category": "Entrega do cronograma",
                      "planned_start": ps, "planned_end": pe, "actual_start": as_, "actual_end": ae, "status": status,
                      "progress": 100 if status == "concluida" else 0, "is_milestone": milestone, "deps": deps, "flags": flags})
    by_area = {}
    for it in items:
        by_area.setdefault(it["area"], {})[it["seq"]] = it
    seen = {}
    for it in items:
        missing = [d for d in it["deps"] if d not in by_area[it["area"]] or d == it["seq"]]
        if missing:
            it["flags"].append(flag("revisar", f"Dependência para item inexistente ({', '.join(map(str, missing))}) — não importada."))
            it["deps"] = [d for d in it["deps"] if d not in missing]
        k = (it["area"], norm(it["title"]))
        if k in seen:
            it["flags"].append(flag("revisar", f"Título repetido na mesma área (item {seen[k]}) — verifique se é duplicidade."))
        seen[k] = it["seq"]
    if future_real:
        notes.append(flag("info", f"{future_real} atividade(s) tinham “tempo real” com datas futuras (cópia do planejado): essas datas reais não foram importadas."))
    notes.append(flag("info", "A situação das atividades foi inferida do “tempo real”: com término real = concluída; só início real = em andamento; sem datas reais = não iniciada."))
    n_ff = 0
    for it in items:
        ff = [d for d in it["deps"] if dep_type(by_area[it["area"]].get(d), it) == "FF"]
        if ff:
            n_ff += len(ff)
            it["flags"].append(flag("info", f"Começa antes do fim do(s) item(ns) {', '.join(map(str, ff))} de que depende: essa(s) dependência(s) entra(m) como Término→Término (confira o tipo)."))
        bad = [d for d in it["deps"] if d not in ff and by_area[it["area"]].get(d) and it["planned_end"] and by_area[it["area"]][d].get("planned_end")
               and (it["planned_start"] or it["planned_end"]) < by_area[it["area"]][d]["planned_end"]]
        if bad:
            it["flags"].append(flag("revisar", f"Datas incompatíveis com a dependência do(s) item(ns) {', '.join(map(str, bad))}: termina antes da predecessora."))
    notes.append(flag("info", "Marcos de entrega: importados apenas onde a planilha tem a célula de marco pintada."))
    notes.append(flag("info", "A planilha só indica “depende de”, sem o tipo. Regra usada: Término→Início quando a sucessora começa depois do fim da predecessora; "
                              f"Término→Término quando começa antes mas termina depois ({n_ff} caso(s)). Sem defasagem. O tipo pode ser alterado depois em cada atividade."))
    return {"activities": items, "notes": notes}


def read_planning(wb, year, today):
    ws = wb.worksheets[0]
    sh = Sheet(ws)
    legend = {}
    for r in range(1, 20):
        t = norm(sh.s(r, 10))
        color = fill_of(ws.cell(r, 13)) or fill_of(ws.cell(r, 10))
        if color and t:
            st = ("em_andamento" if "andamento" in t else "em_revisao" if "melhora" in t else "nao_iniciada" if "nao iniciado" in t else "concluida" if "pronto" in t else None)
            if st:
                legend[fill_of(ws.cell(r, 13)) or color] = st
    items, notes, category, counter = [], [], None, 0
    if not legend:
        notes.append(flag("revisar", "Legenda de cores não encontrada: a situação das demandas não pôde ser lida."))
    for r in range(8, sh.rows + 1):
        b = sh.v(r, 2)
        if not isinstance(b, str) or not b.strip():
            continue
        t = re.sub(r"\s+", " ", b).strip()
        ms = re.match(r"^(\d+)\s*-\s*(.+)$", t)
        if ms and not sh.s(r, 4) and not sh.v(r, 7):
            category = ms.group(2).strip()
            continue
        if norm(t).startswith("atividade realizada") or category is None:
            continue
        counter += 1
        flags = []
        due = sh.v(r, 7)
        pe = due.date().isoformat() if isinstance(due, datetime) else None
        color = fill_of(ws.cell(r, 8))
        status = legend.get(color)
        if not status:
            status = "nao_iniciada"
            flags.append(flag("revisar", "A planilha não indica a situação (célula “Progresso” sem cor): importada como não iniciada."))
        if status == "concluida":
            flags.append(flag("revisar", "Marcada como pronta na planilha, mas sem data real de conclusão: informe a data."))
        if not pe:
            flags.append(flag("info", "Sem prazo na planilha."))
        if norm(t) == "documento":
            flags.append(flag("revisar", "Título genérico (“Documento”): defina um título que identifique a entrega."))
        items.append({"key": f"planejamento:{counter}", "origin": "planejamento", "area": "Gestão do projeto", "seq": None, "title": t,
                      "description": re.sub(r"[ \t]+", " ", sh.s(r, 4)) or None, "category": category, "planned_start": None, "planned_end": pe,
                      "actual_start": None, "actual_end": None, "status": status, "progress": 100 if status == "concluida" else 0,
                      "is_milestone": False, "deps": [], "flags": flags})
    notes.append(flag("info", "A situação de cada demanda foi lida da cor da célula “Progresso”, conforme a legenda da própria planilha."))
    return {"activities": items, "notes": notes}


def read_raci(wb, year, today):
    ws = wb.worksheets[0]
    sh = Sheet(ws)
    people, cols = [], {}
    for c in range(5, sh.cols + 1):
        role_title, name = sh.s(2, c), re.sub(r"\s+", " ", sh.s(4, c))
        if not name or not role_title or norm(role_title) in ("responsavel", "aprovador", "consultado", "informado"):
            continue
        teacher = "tecnico" in norm(role_title)
        key = f"pessoa:{norm(name)}"
        people.append({"key": key, "name": name, "position": role_title.strip(), "area": None if teacher else area_of(role_title), "is_teacher": teacher, "flags": []})
        cols[c] = key
    acts, notes, area = [], [], None
    teacher_keys = {p["key"] for p in people if p["is_teacher"]}
    for r in range(6, sh.rows + 1):
        if sh.s(r, 2):
            area = area_of(sh.s(r, 2)) or sh.s(r, 2)
        title = re.sub(r"\s+", " ", sh.s(r, 3))
        if not title:
            continue
        assign, flags = {}, []
        for c, key in cols.items():
            v = sh.s(r, c).upper()
            if v in ("R", "A", "C", "I"):
                assign[key] = v
            elif v:
                flags.append(flag("revisar", f"Valor “{v}” não é um papel RACI válido — ignorado."))
        approvers = [k for k, v in assign.items() if v == "A"]
        for k in approvers:
            if k in teacher_keys:
                del assign[k]
                flags.append(flag("revisar", "Na planilha o professor técnico estava como aprovador (A): atribuição não importada, pois a regra atual proíbe."))
        if len([k for k, v in assign.items() if v == "A"]) > 1:
            for k in [k for k, v in assign.items() if v == "A"][1:]:
                del assign[k]
            flags.append(flag("revisar", "Mais de um aprovador (A) na planilha: apenas o primeiro foi mantido."))
        if "A" not in assign.values():
            flags.append(flag("revisar", "Sem aprovador final (A)."))
        rs = [k for k, v in assign.items() if v == "R"]
        if not rs:
            flags.append(flag("revisar", "Sem responsável pela execução (R)."))
        elif len(rs) > 1:
            flags.append(flag("info", "Mais de um responsável (R): o responsável principal ficou em branco para a equipe definir."))
        flags.append(flag("info", "A planilha RACI não traz prazo para esta entrega."))
        acts.append({"key": f"raci:{norm(area)}:{norm(title)}", "origin": "raci", "area": area, "seq": None, "title": title, "category": "Entregável (RACI)",
                     "planned_start": None, "planned_end": None, "actual_start": None, "actual_end": None, "status": "nao_iniciada", "progress": 0,
                     "is_milestone": False, "deps": [], "flags": flags, "raci": assign, "owner_key": rs[0] if len(rs) == 1 else None})
    notes.append(flag("info", "As 18 entregas da matriz RACI viram atividades (categoria “Entregável (RACI)”) para que a matriz fique ligada ao planejamento. "
                              "Elas não têm datas na planilha." if len(acts) == 18 else
                      "As entregas da matriz RACI viram atividades (categoria “Entregável (RACI)”), sem datas."))
    return {"people": people, "activities": acts, "notes": notes}


def read_risks(wb, year, today):
    ws = wb.worksheets[0]
    sh = Sheet(ws)
    header, cols = None, {}
    for r in range(1, min(sh.rows, 40) + 1):
        row = {norm(sh.s(r, c)): c for c in range(1, sh.cols + 1) if sh.s(r, c)}
        if "risco" in row and "probabilidade" in row and "impacto" in row:
            header, cols = r, row
            break
    if not header:
        return {"risks": [], "notes": [flag("revisar", "Cabeçalho da matriz de riscos não encontrado.")]}

    def col(*names):
        for n in names:
            for k, c in cols.items():
                if k.startswith(n):
                    return c
    c_r, c_c, c_q, c_p, c_i, c_n, c_m, c_o = col("risco"), col("causa"), col("consequencia"), col("probabilidade"), col("impacto"), col("nivel"), col("plano"), col("responsavel")
    risks, n = [], 0
    for r in range(header + 1, sh.rows + 1):
        title = re.sub(r"\s+", " ", sh.s(r, c_r))
        if not title:
            continue
        n += 1
        flags = []
        p_txt, i_txt, lvl = sh.s(r, c_p), sh.s(r, c_i), sh.s(r, c_n) if c_n else ""
        p, i = PROB_MAP.get(norm(p_txt)), IMPACT_MAP.get(norm(i_txt))
        if not p:
            flags.append(flag("revisar", f"Probabilidade “{p_txt}” não corresponde à escala — deixada em branco."))
        if not i:
            flags.append(flag("revisar", f"Impacto “{i_txt}” não corresponde à escala — deixado em branco."))
        owner_txt = re.sub(r"\s+", " ", sh.s(r, c_o)) if c_o else ""
        areas = []
        for part in re.split(r"\+|/| e ", owner_txt):
            a = area_of(part)
            if a and a not in areas:
                areas.append(a)
            elif part.strip() and not a:
                flags.append(flag("revisar", f"Responsável “{part.strip()}” não corresponde a uma área cadastrada."))
        risks.append({"key": f"risco:{norm(title)}", "code": f"R-{n:03d}", "title": title, "cause": sh.s(r, c_c) or None, "consequence": sh.s(r, c_q) or None,
                      "probability": p, "impact": i, "original_probability": p_txt or None, "original_impact": i_txt or None, "original_level": lvl or None,
                      "mitigation": sh.s(r, c_m) or None, "owner_text": owner_txt or None, "areas": areas, "flags": flags})
    notes = [flag("info", "Regra de conversão da escala textual para 1–5 — Probabilidade: Raro=1, Improvável=2, Possível=3, Provável=4, Muito provável=5; "
                          "Impacto: Muito baixo=1, Baixo=2, Médio=3, Alto=4, Muito alto=5. O nível original da planilha é preservado ao lado do nível calculado (P×I)."),
             flag("info", "A coluna “Responsável” da planilha indica áreas, não pessoas: as áreas foram vinculadas e o responsável (pessoa) ficou em branco.")]
    return {"risks": risks, "notes": notes}


_CODE_RX = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s+(\S.*?)\s*$")


def read_wbs(wb, year, today):
    seen, items, dup, notes = {}, [], 0, []
    for si, ws in enumerate(wb.worksheets):
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                lines = [ln for ln in c.value.split("\n") if ln.strip()]
                multi = len(lines) > 1
                for ln in lines:
                    m = _CODE_RX.match(ln)
                    if not m:
                        continue
                    code, name = m.group(1), re.sub(r"\s+", " ", m.group(2)).strip()
                    if code in seen:
                        if norm(seen[code]["name"]) == norm(name):
                            dup += 1
                        else:
                            seen[code]["flags"].append(flag("revisar", f"O código {code} aparece também como “{name}” ({ws.title}!{c.coordinate})."))
                        continue
                    it = {"key": f"eap:{code}", "code": code, "name": name, "flags": []}
                    if multi:
                        it["flags"].append(flag("info", f"Lido de uma célula que reúne vários itens ({ws.title}!{c.coordinate})."))
                    seen[code] = it
                    items.append(it)
    items.sort(key=lambda x: tuple(int(p) for p in x["code"].split(".")))
    for it in items:
        parts = it["code"].split(".")
        it["parent_code"] = ".".join(parts[:-1]) if len(parts) > 1 else None
        if it["parent_code"] and it["parent_code"] not in seen:
            it["flags"].append(flag("revisar", f"Elemento superior {it['parent_code']} não encontrado — importado no nível raiz."))
            it["parent_code"] = None
        top = seen.get(".".join(parts[:2])) if len(parts) >= 2 else None
        it["area"] = area_of(top["name"]) if top else None
    if dup:
        notes.append(flag("info", f"{dup} célula(s) repetiam itens já lidos (mesmo código e nome), inclusive na segunda aba: contadas como duplicadas e ignoradas."))
    return {"wbs": items, "notes": notes, "signature": hashlib.sha256(json.dumps([(i["code"], i["name"]) for i in items]).encode()).hexdigest()}


def read_budget(wb, year, today):
    ws = wb.worksheets[0]
    sh = Sheet(ws)
    header = next((r for r in range(1, 30) if "custos do projeto" in norm(sh.s(r, 2))), None)
    if not header:
        return {"expenses": [], "notes": [flag("revisar", "Tabela de custos não encontrada.")]}
    cmap = {}
    for c in range(2, 21):
        t = norm(sh.s(header, c))
        if t.startswith("preco"):
            cmap["unit"] = c
        elif t.startswith("quantidade"):
            cmap["qty"] = c
        elif t == "custo":
            cmap["cost"] = c
        elif t.startswith("adquirido"):
            cmap["sponsored"] = c
    right = next((c for c in range(21, sh.cols + 1) if "custos do projeto" in norm(sh.s(header, c))), None)
    rmap = {}
    if right:
        for c in range(right, sh.cols + 1):
            t = norm(sh.s(header, c))
            if t.startswith("preco"):
                rmap["unit"] = c
            elif t.startswith("quantidade"):
                rmap["qty"] = c
            elif t == "custo":
                rmap["cost"] = c
    items, notes, area, sums, totals = [], [], None, {}, {}
    summary = {}
    for r in range(header + 1, sh.rows + 1):
        name = re.sub(r"\s+", " ", sh.s(r, 2))
        if name:
            mt = re.match(r"^total estimado:?\s*(?:r\$)?\s*([\d\.,]*)", name, re.I)
            if mt:
                raw = mt.group(1).strip()
                if raw and area:
                    try:
                        totals[area] = float(raw.replace(".", "").replace(",", ".")) if "," in raw else float(raw)
                    except ValueError:
                        pass
                continue
            if area_of(name.rstrip("- ")) and not isinstance(sh.v(r, cmap.get("unit", 8)), (int, float)):
                area = area_of(name.rstrip("- "))
                continue
            unit = sh.v(r, cmap.get("unit", 8))
            if isinstance(unit, (int, float)) and area:
                flags = []
                qraw = sh.s(r, cmap.get("qty", 12))
                mq = re.match(r"^\s*([\d\.,]+)\s*(.*)$", qraw)
                qty, unit_label = (float(mq.group(1).replace(",", ".")), mq.group(2).strip() or None) if mq else (None, re.sub(r"^[xX]\s*", "", qraw).strip() or None)
                if qty is None and qraw:
                    flags.append(flag("info", f"Quantidade não definida na planilha (“{qraw}”)."))
                cost = sh.v(r, cmap.get("cost", 15))
                cost = float(cost) if isinstance(cost, (int, float)) else None
                sponsored = norm(sh.s(r, cmap.get("sponsored", 18))) == "sim"
                if qty is not None and cost is not None and abs(qty * float(unit) - cost) > 0.009:
                    flags.append(flag("revisar", f"Preço unitário × quantidade ({float(unit):.2f} × {qty:g} = {qty * float(unit):.2f}) é diferente do custo informado ({cost:.2f}). Os valores foram importados como estão."))
                sponsor = None
                ms = re.search(r"\(([^)]+)\)", name)
                if ms and sponsored:
                    sponsor = ms.group(1).strip()
                if sponsored and not cost:
                    flags.append(flag("info", "Item obtido por patrocínio sem valor estimado: o apoio não entra nos totais até que um valor seja informado."))
                actual = None
                if rmap.get("cost"):
                    av = sh.v(r, rmap["cost"])
                    actual = float(av) if isinstance(av, (int, float)) else None
                items.append({"key": f"orcamento:{norm(area)}:{norm(name)}", "name": name, "area": area, "quantity": qty, "unit": unit_label,
                              "unit_planned": float(unit), "total_planned": cost if cost is not None else (qty * float(unit) if qty is not None else None),
                              "total_actual": actual, "sponsored": sponsored, "sponsor_name": sponsor, "flags": flags})
                if not sponsored:
                    sums[area] = round(sums.get(area, 0) + (cost or 0), 2)
        for c in range(1, sh.cols + 1):
            t = norm(sh.s(r, c))
            if t in ("valor estipulado para 2026", "valor a arrecadar", "valor ja arrecadado", "valor restante") or t.startswith("valor estipulado"):
                for rr in range(r + 1, min(r + 4, sh.rows + 1)):
                    val = sh.v(rr, c)
                    if isinstance(val, (int, float)):
                        summary["budget" if t.startswith("valor estipulado") else "goal" if t == "valor a arrecadar" else "raised" if t == "valor ja arrecadado" else "remaining"] = float(val)
                        break
    checks = []
    for a, tot in totals.items():
        checks.append({"check": f"Total estimado — {a}", "expected": tot, "actual": sums.get(a, 0), "ok": abs(tot - sums.get(a, 0)) < 0.01})
    if right and not any(i["total_actual"] is not None for i in items):
        notes.append(flag("info", "A tabela “Controle orçamentário final” está sem valores: nenhum valor realizado foi importado."))
    if summary.get("raised"):
        notes.append(flag("revisar", f"A planilha indica R$ {summary['raised']:.2f} já arrecadados, sem data nem comprovante. O valor NÃO foi lançado no caixa: "
                                     "registre o recebimento na receita correspondente quando confirmar que o dinheiro entrou."))
    return {"expenses": items, "budget": summary, "checks": checks, "notes": notes}


def read_stakeholders(wb, year, today):
    sh = Sheet(wb.worksheets[0])
    header = next((r for r in range(1, 30) if "patrocinadores" in norm(sh.s(r, 2))), None)
    if not header:
        return {"stakeholders": [], "notes": [flag("revisar", "Tabela de stakeholders não encontrada.")]}
    items = []
    for r in range(header + 1, sh.rows + 1):
        name = re.sub(r"\s+", " ", sh.s(r, 2))
        if not name:
            continue
        flags = []
        need, val, conf = sh.s(r, 4) or None, sh.v(r, 7), norm(sh.s(r, 8))
        amount = float(val) if isinstance(val, (int, float)) else None
        status = "confirmado" if conf == "sim" else "recusado" if conf in ("nao", "não") else "prospectado"
        if not need and amount is None and not conf:
            flags.append(flag("revisar", "Registro apenas com o nome na planilha: complete os dados ou exclua."))
        elif not conf:
            flags.append(flag("revisar", "Sem resposta de confirmação na planilha: importado como prospectado."))
        if amount is not None and status == "confirmado":
            flags.append(flag("revisar", f"Valor confirmado de R$ {amount:.2f}: foi criada uma receita “a receber”. Registre o recebimento com a data real se o dinheiro já entrou."))
        items.append({"key": f"stakeholder:{norm(name)}", "name": name, "need": need, "confirmed": amount if status == "confirmado" else None,
                      "requested": None, "status": status, "contribution_type": "Financeira" if amount is not None else None, "flags": flags})
    return {"stakeholders": items, "notes": [flag("info", "“Confirmação = SIM” foi importada como negociação confirmada. Valores com traços (“-----”) ficaram em branco.")]}


def read_scope(wb, year, today):
    sh = Sheet(wb.worksheets[0])
    data = {"deliverables": [], "specific_objectives": []}
    flags = []
    for r in range(1, sh.rows + 1):
        for c in range(1, sh.cols + 1):
            label = norm(sh.s(r, c)).rstrip(" :")
            if not label or not isinstance(sh.v(r, c), str):
                continue
            val = None
            for cc in range(c + 1, min(c + 6, sh.cols + 1)):
                if sh.v(r, cc) is not None and txt(sh.v(r, cc)):
                    val = sh.v(r, cc)
                    break
            if val is None:
                continue
            if label.startswith("gestora do projeto") or label.startswith("gestor do projeto"):
                data["manager_name"] = txt(val)
            elif label == "data":
                data["date"] = val.date().isoformat() if isinstance(val, datetime) else None
            elif label == "nome do projeto":
                data["project_name"] = txt(val)
            elif label == "objetivo do projeto":
                data["general_objective"] = txt(val)
            elif label.startswith("criterios de aprovacao"):
                data["approval_criteria"] = txt(val)
            elif label == "riscos":
                data["main_risks"] = txt(val)
            elif area_of(label) and c <= 3:
                for part in txt(val).rstrip(".").split(","):
                    p = part.strip()
                    if p:
                        data["deliverables"].append({"text": p[0].upper() + p[1:], "area": area_of(label), "main": False})
    if not data.get("project_name"):
        flags.append(flag("revisar", "Nome do projeto não encontrado."))
    flags.append(flag("info", "A planilha não traz: objetivos específicos, justificativa, premissas, restrições, exclusões e critérios de aceitação — esses campos ficam em branco para a equipe preencher."))
    return {"scope": {"key": "escopo", "data": data, "flags": flags}, "notes": []}


def read_evidence(wb, year, today):
    sh = Sheet(wb.worksheets[0])
    month, week_cols, entries, notes = None, {}, [], []
    for r in range(1, sh.rows + 1):
        label = norm(sh.s(r, 2)).rstrip(" :")
        if label == "mes":
            month = MONTH_NUM.get(norm(sh.s(r, 3)))
        elif label == "semana":
            for c in range(3, sh.cols + 1):
                v = sh.v(r, c)
                if isinstance(v, (int, float)) and 1 <= int(v) <= 6:
                    week_cols[c] = int(v)
        elif area_of(label) and week_cols:
            for c, wk in week_cols.items():
                t = sh.s(r, c)
                if t:
                    entries.append({"week": wk, "area": area_of(label), "text": t})
    if not month:
        return {"evidence": None, "notes": [flag("revisar", "Mês de referência não encontrado na planilha de evidências.")]}
    weeks = []
    for wk in sorted(set(week_cols.values())):
        es = [e for e in entries if e["week"] == wk]
        weeks.append({"key": f"evidencia:{year}-{month:02d}:{wk}", "month": f"{year}-{month:02d}", "week_number": wk, "entries": es,
                      "flags": [flag("revisar", "A planilha não informa as datas de início e término da semana nem o responsável pelo acompanhamento.")] +
                               ([flag("info", f"Apenas {len(es)} área(s) têm registro nesta semana.")] if 0 < len(es) < 4 else []) +
                               ([flag("info", "Semana sem nenhum registro na planilha — não será criada.")] if not es else [])})
    notes.append(flag("info", f"O ano do mês de referência não consta na planilha; foi usado o ano da temporada ({year}). Os textos foram importados no campo “Atividades realizadas”."))
    return {"evidence": {"weeks": weeks}, "notes": notes}


READERS = {"gantt": read_gantt, "planejamento": read_planning, "raci": read_raci, "riscos": read_risks, "eap": read_wbs,
           "orcamento": read_budget, "stakeholders": read_stakeholders, "escopo": read_scope, "evidencias": read_evidence}


# ----------------------------------------------------------------------------- montagem da pré-visualização
def build_staging(db, season, uploads):
    """Lê os arquivos e monta a pré-visualização. `uploads` = [{name, data}]. Nada é gravado nas tabelas de negócio."""
    today = today_local(db)
    year = season["year"]
    st = {"files": [], "people": [], "wbs": [], "activities": [], "risks": [], "expenses": [], "stakeholders": [], "scope": None,
          "evidence": [], "budget": {}, "checks": [], "notes": []}
    wbs_sig = None
    seen_hash = {}
    for up in uploads:
        info = {"name": up["name"], "sha256": hashlib.sha256(up["data"]).hexdigest(), "size": len(up["data"]), "kind": None, "status": "lido", "message": None}
        st["files"].append(info)
        if info["sha256"] in seen_hash:
            info.update(status="duplicado", message=f"Arquivo idêntico a “{seen_hash[info['sha256']]}”.")
            continue
        seen_hash[info["sha256"]] = up["name"]
        try:
            wb = openpyxl.load_workbook(io.BytesIO(up["data"]), data_only=False)
        except Exception as exc:
            info.update(status="erro", message=f"Não foi possível abrir a planilha ({type(exc).__name__}).")
            continue
        kind = detect(wb)
        info["kind"] = kind
        info["sheets"] = len(wb.worksheets)
        if not kind:
            info.update(status="ignorado", message="Planilha não reconhecida — nenhum dado importado.")
            continue
        try:
            res = READERS[kind](wb, year, today)
        except Exception as exc:
            info.update(status="erro", message=f"Falha ao ler a planilha ({type(exc).__name__}: {exc}).")
            continue
        if kind == "eap":
            if wbs_sig and res["signature"] == wbs_sig["sig"]:
                info.update(status="duplicado", message=f"Conteúdo idêntico ao de “{wbs_sig['name']}” ({len(res['wbs'])} elementos): tratado como a mesma versão e não importado em duplicidade.")
                continue
            if wbs_sig:
                info.update(status="revisar", message=f"Versão diferente de “{wbs_sig['name']}”: apenas a primeira foi considerada. Revise as diferenças antes de importar esta.")
                continue
            wbs_sig = {"sig": res["signature"], "name": up["name"]}
        for n in res.get("notes", []):
            st["notes"].append({**n, "file": up["name"]})
        for sec in ("people", "wbs", "activities", "risks", "expenses", "stakeholders"):
            for it in res.get(sec, []):
                it["file"] = up["name"]
                st[sec].append(it)
        if res.get("scope"):
            st["scope"] = {**res["scope"], "file": up["name"]}
        if res.get("evidence"):
            for w in res["evidence"]["weeks"]:
                w["file"] = up["name"]
                st["evidence"].append(w)
        if res.get("budget"):
            st["budget"] = res["budget"]
        st["checks"] += res.get("checks", [])
        info["message"] = KIND_LABEL[kind]
    # possíveis duplicatas entre Planejamento e Gantt (mesma área)
    gantt = [a for a in st["activities"] if a["origin"] == "gantt"]
    for a in st["activities"]:
        a["action"] = "importar"
        if a["origin"] == "planejamento":
            tk = tokens(a["title"])
            match = next((g for g in gantt if g["area"] == a["area"] and tokens(g["title"]) == tk), None)
            if match:
                a["merge_into"] = match["key"]
                a["merge_title"] = f"{match['seq']}. {match['title']}"
                a["action"] = "mesclar"
                a["flags"].append(flag("revisar", f"Parece ser a mesma entrega do cronograma “{match['seq']}. {match['title']}”. Padrão: mesclar (não cria atividade nova)."))
    # marca o que já existe no banco (reimportação não duplica)
    existing = {
        "activities": {r["source"] for r in db.q("SELECT source FROM activities WHERE season_id=? AND source IS NOT NULL AND deleted_at IS NULL", (season["id"],))},
        "wbs": {r["source"] for r in db.q("SELECT source FROM wbs_elements WHERE season_id=? AND source IS NOT NULL AND deleted_at IS NULL", (season["id"],))},
        "risks": {r["source"] for r in db.q("SELECT source FROM risks WHERE season_id=? AND source IS NOT NULL AND deleted_at IS NULL", (season["id"],))},
        "expenses": {r["source"] for r in db.q("SELECT source FROM expenses WHERE season_id=? AND source IS NOT NULL AND deleted_at IS NULL", (season["id"],))},
        "stakeholders": {r["source"] for r in db.q("SELECT source FROM stakeholders WHERE season_id=? AND source IS NOT NULL AND deleted_at IS NULL", (season["id"],))},
        "evidence": {r["source"] for r in db.q("SELECT source FROM evidence_weeks WHERE season_id=? AND source IS NOT NULL AND deleted_at IS NULL", (season["id"],))},
    }
    users = db.q("SELECT id, name FROM users WHERE deleted_at IS NULL")
    for p in st["people"]:
        p["action"] = "importar"
        m = next((u for u in users if tokens(u["name"]) == tokens(p["name"]) or (len(tokens(u["name"])) >= 2 and (tokens(u["name"]) <= tokens(p["name"]) or tokens(p["name"]) <= tokens(u["name"])))), None)
        if m:
            p["match_user_id"], p["exists"] = m["id"], True
            p["flags"].append(flag("info", f"Já existe o usuário “{m['name']}”: será vinculado (não duplica)."))
    for sec in ("wbs", "risks", "expenses", "stakeholders", "evidence"):
        for it in st[sec]:
            it.setdefault("action", "importar")
            if "planilha:" + it["key"] in existing[sec]:
                it["exists"] = True
    for a in st["activities"]:
        if "planilha:" + a["key"] in existing["activities"]:
            a["exists"] = True
    for w in st["evidence"]:
        if not w["entries"]:
            w["action"] = "ignorar"
    if st["scope"]:
        st["scope"]["action"] = "importar"
        if db.one("SELECT 1 FROM scopes WHERE season_id=?", (season["id"],)):
            st["scope"]["exists"] = True
            st["scope"]["flags"].append(flag("info", "Já existe um escopo nesta temporada: o da planilha não substituirá o atual."))
    # conferências cruzadas
    conf_total = round(sum(s["confirmed"] or 0 for s in st["stakeholders"]), 2)
    if st["budget"].get("raised") is not None and st["stakeholders"]:
        st["checks"].append({"check": "Valor já arrecadado (orçamento) × soma dos valores confirmados (stakeholders)", "expected": st["budget"]["raised"], "actual": conf_total,
                             "ok": abs(st["budget"]["raised"] - conf_total) < 0.01})
    if {"goal", "raised", "remaining"} <= set(st["budget"]):
        b = st["budget"]
        st["checks"].append({"check": "Valor restante = valor a arrecadar − valor já arrecadado", "expected": round(b["goal"] - b["raised"], 2), "actual": b["remaining"],
                             "ok": abs(b["goal"] - b["raised"] - b["remaining"]) < 0.01})
    names = {norm(s["name"]) for s in st["stakeholders"]} | {norm(r["name"]) for r in db.q("SELECT name FROM stakeholders WHERE season_id=? AND deleted_at IS NULL", (season["id"],))}
    for e in st["expenses"]:
        if e.get("sponsor_name") and norm(e["sponsor_name"]) not in names:
            e["flags"].append(flag("revisar", f"Patrocinador “{e['sponsor_name']}” não encontrado na lista de stakeholders: item importado sem vínculo."))
    st["summary"] = summarize(st)
    return st


def _needs_review(it):
    return any(f["level"] == "revisar" for f in it.get("flags", []))


def summarize(st):
    out = {}
    for sec in ("people", "wbs", "activities", "risks", "expenses", "stakeholders", "evidence"):
        items = st[sec]
        out[sec] = {"label": SECTION_LABEL[sec], "total": len(items),
                    "importar": sum(1 for i in items if i.get("action") == "importar" and not i.get("exists")),
                    "mesclar": sum(1 for i in items if i.get("action") == "mesclar" and not i.get("exists")),
                    "ignorar": sum(1 for i in items if i.get("action") == "ignorar" and not i.get("exists")),
                    "duplicados": sum(1 for i in items if i.get("exists")),
                    "revisar": sum(1 for i in items if _needs_review(i) and i.get("action") != "ignorar" and not i.get("exists"))}
    if st.get("scope"):
        s = st["scope"]
        out["scope"] = {"label": "Escopo", "total": 1, "importar": int(s.get("action") == "importar" and not s.get("exists")), "mesclar": 0,
                        "ignorar": int(s.get("action") == "ignorar"), "duplicados": int(bool(s.get("exists"))), "revisar": 0}
    out["dependencies"] = {"label": "Dependências do cronograma", "total": sum(len(a["deps"]) for a in st["activities"] if a.get("action") == "importar")}
    out["raci"] = {"label": "Atribuições RACI", "total": sum(len(a.get("raci") or {}) for a in st["activities"] if a.get("action") == "importar")}
    out["evidence_entries"] = {"label": "Registros de evidência por área", "total": sum(len(w["entries"]) for w in st["evidence"] if w.get("action") == "importar")}
    return out


def dep_type(pred, succ):
    """A planilha só diz "depende de". Usa Término→Início quando as datas planejadas permitem; se a sucessora começa antes
    do fim da predecessora mas termina depois dela (ex.: "100% do portfólio" depende dos relatórios), usa Término→Término."""
    if not pred or not pred.get("planned_end") or not succ.get("planned_end"):
        return "FS"
    s_start = succ.get("planned_start") or succ["planned_end"]
    if s_start >= pred["planned_end"]:
        return "FS"
    return "FF" if succ["planned_end"] >= pred["planned_end"] else "FS"


# ----------------------------------------------------------------------------- gravação
def apply_import(ctx, imp, st):
    """Grava a importação confirmada. Executa dentro da transação da requisição: ou entra tudo, ou nada."""
    db = ctx.db
    sid = imp["season_id"]
    ts = now_iso()
    uid = ctx.uid
    res = {k: {"importados": 0, "mesclados": 0, "ignorados": 0, "duplicados": 0, "para_revisar": 0} for k in
           ("people", "wbs", "activities", "dependencies", "raci", "risks", "expenses", "stakeholders", "scope", "evidence", "evidence_entries")}
    areas = {a["name"]: a["id"] for a in db.q("SELECT id, name FROM areas WHERE deleted_at IS NULL")}

    def area_id(name):
        if not name:
            return None
        if name not in areas:
            aid = new_id()
            db.insert("areas", {"id": aid, "name": name, "sort": len(areas) + 1, "created_at": ts, "updated_at": ts, "created_by": uid})
            areas[name] = aid
        return areas[name]

    def review(it):
        notes = [f["text"] for f in it.get("flags", []) if f["level"] == "revisar"]
        return (1, " ".join(notes)[:2000]) if notes else (0, None)

    def skip(sec, it):
        if it.get("exists"):
            res[sec]["duplicados"] += 1
            return True
        if it.get("action") == "ignorar":
            res[sec]["ignorados"] += 1
            return True
        return False
    # pessoas
    pmap = {}
    for p in st["people"]:
        if p.get("match_user_id"):
            pmap[p["key"]] = p["match_user_id"]
            u = db.one("SELECT * FROM users WHERE id=?", (p["match_user_id"],))
            db.ex("UPDATE users SET area_id=COALESCE(area_id, ?), position=COALESCE(position, ?), is_teacher=CASE WHEN is_teacher=1 OR ?=1 THEN 1 ELSE 0 END, updated_at=?, version=version+1 WHERE id=?",
                  (area_id(p["area"]), p["position"], 1 if p["is_teacher"] else 0, ts, u["id"]))
            res["people"]["duplicados"] += 1
            continue
        if p.get("action") == "ignorar":
            res["people"]["ignorados"] += 1
            continue
        nid = new_id()
        db.insert("users", {"id": nid, "name": p["name"], "role": "membro", "is_teacher": 1 if p["is_teacher"] else 0, "area_id": area_id(p["area"]),
                            "position": p["position"], "created_at": ts, "updated_at": ts, "created_by": uid})
        pmap[p["key"]] = nid
        res["people"]["importados"] += 1
    # EAP
    wmap = {r["source"][len("planilha:eap:"):]: r["id"] for r in db.q("SELECT id, source FROM wbs_elements WHERE season_id=? AND source LIKE 'planilha:eap:%' AND deleted_at IS NULL", (sid,))}
    for w in st["wbs"]:
        if skip("wbs", w):
            continue
        nid = new_id()
        nr, note = review(w)
        db.insert("wbs_elements", {"id": nid, "season_id": sid, "parent_id": wmap.get(w["parent_code"]), "sort": int(w["code"].split(".")[-1]), "name": w["name"],
                                   "area_id": area_id(w["area"]), "needs_review": nr, "review_note": note, "source": "planilha:" + w["key"],
                                   "created_at": ts, "updated_at": ts, "created_by": uid})
        wmap[w["code"]] = nid
        res["wbs"]["importados"] += 1
        res["wbs"]["para_revisar"] += nr
    # stakeholders (antes do orçamento, para vincular itens patrocinados)
    from .stakeholders import ST, sync_income
    smap = {norm(r["name"]): r["id"] for r in db.q("SELECT id, name FROM stakeholders WHERE season_id=? AND deleted_at IS NULL", (sid,))}
    for s in st["stakeholders"]:
        if skip("stakeholders", s):
            continue
        nid = new_id()
        nr, note = review(s)
        from ..core import to_cents
        db.insert("stakeholders", {"id": nid, "season_id": sid, "name": s["name"], "type": "Patrocinador", "need": s["need"], "contribution_type": s["contribution_type"],
                                   "confirmed_cents": to_cents(s["confirmed"]), "status": s["status"], "needs_review": nr, "review_note": note,
                                   "source": "planilha:" + s["key"], "created_at": ts, "updated_at": ts, "created_by": uid})
        smap[norm(s["name"])] = nid
        sync_income(ctx, db.one("SELECT * FROM stakeholders WHERE id=?", (nid,)))
        db.ex("UPDATE incomes SET needs_review=1, review_note=? WHERE stakeholder_id=? AND import_source='stakeholder:auto'",
              ("Valor confirmado na planilha, sem data de recebimento: registre o recebimento quando o dinheiro entrar.", nid))
        res["stakeholders"]["importados"] += 1
        res["stakeholders"]["para_revisar"] += nr
    # orçamento
    from ..core import to_cents
    for e in st["expenses"]:
        if skip("expenses", e):
            continue
        nr, note = review(e)
        sponsor_id = smap.get(norm(e["sponsor_name"])) if e.get("sponsor_name") else None
        db.insert("expenses", {"id": new_id(), "season_id": sid, "name": e["name"], "area_id": area_id(e["area"]), "quantity": e["quantity"], "unit": e["unit"],
                               "unit_planned_cents": to_cents(e["unit_planned"]), "total_planned_cents": to_cents(e["total_planned"]),
                               "total_actual_cents": to_cents(e["total_actual"]), "acquisition_type": "item_patrocinado" if e["sponsored"] else "proprio",
                               "is_cash": 0 if e["sponsored"] else 1, "stakeholder_id": sponsor_id, "status": "planejada",
                               "funding_source": "Patrocínio" if e["sponsored"] else None, "needs_review": nr, "review_note": note,
                               "source": "planilha:" + e["key"], "created_at": ts, "updated_at": ts, "created_by": uid})
        if sponsor_id:
            db.ex("UPDATE stakeholders SET contribution_type=COALESCE(contribution_type, 'Serviço') WHERE id=?", (sponsor_id,))
        res["expenses"]["importados"] += 1
        res["expenses"]["para_revisar"] += nr
    b = st.get("budget") or {}
    season = db.one("SELECT * FROM seasons WHERE id=?", (sid,))
    if b.get("budget") and not season["budget_planned_cents"]:
        db.ex("UPDATE seasons SET budget_planned_cents=?, updated_at=?, version=version+1 WHERE id=?", (to_cents(b["budget"]), ts, sid))
    if b.get("goal") and not season["fundraising_goal_cents"]:
        db.ex("UPDATE seasons SET fundraising_goal_cents=?, updated_at=?, version=version+1 WHERE id=?", (to_cents(b["goal"]), ts, sid))
    # atividades
    amap = {r["source"][len("planilha:"):]: r["id"] for r in db.q("SELECT id, source FROM activities WHERE season_id=? AND source LIKE 'planilha:%' AND deleted_at IS NULL", (sid,))}
    by_key = {a["key"]: a for a in st["activities"]}
    order = sorted(st["activities"], key=lambda a: {"gantt": 0, "raci": 1, "planejamento": 2}[a["origin"]])
    for a in order:
        if skip("activities", a):
            continue
        if a.get("action") == "mesclar" and a.get("merge_into") and amap.get(a["merge_into"]):
            tgt = amap[a["merge_into"]]
            extra = f"Planejamento de gestão — prazo: {fmt_date(a['planned_end'])}." if a.get("planned_end") else None
            if extra:
                db.ex("UPDATE activities SET description=COALESCE(description, ?), notes=CASE WHEN notes IS NULL OR notes='' THEN ? ELSE notes || ? END, updated_at=? WHERE id=?",
                      (a.get("description"), extra, "\n" + extra, ts, tgt))
            else:
                db.ex("UPDATE activities SET description=COALESCE(description, ?), updated_at=? WHERE id=?", (a.get("description"), ts, tgt))
            amap[a["key"]] = tgt
            res["activities"]["mesclados"] += 1
            continue
        nid = new_id()
        nr, note = review(a)
        owner = pmap.get(a.get("owner_key")) if a.get("owner_key") else None
        db.insert("activities", {"id": nid, "season_id": sid, "area_id": area_id(a["area"]), "seq": a["seq"], "title": a["title"][:200], "description": a.get("description"),
                                 "category": a.get("category"), "owner_id": owner, "status": a["status"], "progress": a["progress"],
                                 "planned_start": a["planned_start"], "planned_end": a["planned_end"], "baseline_start": a["planned_start"], "baseline_end": a["planned_end"],
                                 "actual_start": a["actual_start"], "actual_end": a["actual_end"], "is_milestone": 1 if a["is_milestone"] else 0,
                                 "approval_status": "em_execucao" if a["status"] in ("em_andamento", "em_revisao") else "rascunho",
                                 "needs_review": nr, "review_note": note, "source": "planilha:" + a["key"],
                                 "created_at": ts, "updated_at": ts, "created_by": uid})
        amap[a["key"]] = nid
        res["activities"]["importados"] += 1
        res["activities"]["para_revisar"] += nr
        for pk, role in (a.get("raci") or {}).items():
            if pk in pmap:
                db.ex("INSERT OR IGNORE INTO raci (activity_id, user_id, role) VALUES (?,?,?)", (nid, pmap[pk], role))
                res["raci"]["importados"] += 1
    from .activities import would_cycle
    for a in st["activities"]:
        if a["origin"] != "gantt" or a.get("action") != "importar" or a.get("exists"):
            continue
        for dseq in a["deps"]:
            pk = f"gantt:{norm(a['area'])}:{dseq}"
            pred, succ = amap.get(pk), amap.get(a["key"])
            if not pred or not succ or (by_key.get(pk, {}).get("action") == "ignorar"):
                res["dependencies"]["ignorados"] += 1
                continue
            if db.one("SELECT 1 FROM dependencies WHERE predecessor_id=? AND successor_id=?", (pred, succ)):
                res["dependencies"]["duplicados"] += 1
                continue
            if would_cycle(db, pred, succ):
                res["dependencies"]["ignorados"] += 1
                db.ex("UPDATE activities SET needs_review=1, review_note=COALESCE(review_note || ' ', '') || ? WHERE id=?",
                      (f"Dependência do item {dseq} não importada: formaria um ciclo.", succ))
                continue
            db.insert("dependencies", {"id": new_id(), "predecessor_id": pred, "successor_id": succ, "type": dep_type(by_key.get(pk), a), "lag_days": 0, "created_by": uid, "created_at": ts})
            res["dependencies"]["importados"] += 1
    # riscos
    for r in st["risks"]:
        if skip("risks", r):
            continue
        nid = new_id()
        nr, note = review(r)
        code = r["code"]
        while db.one("SELECT 1 FROM risks WHERE season_id=? AND code=?", (sid, code)):
            code = f"R-{int(code[2:]) + 1:03d}"
        db.insert("risks", {"id": nid, "season_id": sid, "code": code, "title": r["title"], "cause": r["cause"], "consequence": r["consequence"],
                            "probability": r["probability"], "impact": r["impact"], "original_probability": r["original_probability"],
                            "original_impact": r["original_impact"], "original_level": r["original_level"], "mitigation": r["mitigation"],
                            "owner_text": r["owner_text"], "status": "aberto", "needs_review": nr, "review_note": note, "source": "planilha:" + r["key"],
                            "created_at": ts, "updated_at": ts, "created_by": uid})
        for an in r["areas"]:
            db.ex("INSERT OR IGNORE INTO risk_areas (risk_id, area_id) VALUES (?,?)", (nid, area_id(an)))
        if r["probability"] and r["impact"]:
            db.insert("risk_reviews", {"id": new_id(), "risk_id": nid, "at": ts, "user_id": uid, "probability": r["probability"], "impact": r["impact"],
                                       "note": f"Importado da planilha (probabilidade “{r['original_probability']}”, impacto “{r['original_impact']}”, nível “{r['original_level']}”)."})
        res["risks"]["importados"] += 1
        res["risks"]["para_revisar"] += nr
    # escopo
    sc = st.get("scope")
    if sc:
        if sc.get("exists"):
            res["scope"]["duplicados"] += 1
        elif sc.get("action") == "ignorar":
            res["scope"]["ignorados"] += 1
        else:
            from .scope import clean_scope
            data = dict(sc["data"])
            data["deliverables"] = [{"text": x["text"], "area_id": area_id(x["area"]), "main": x["main"]} for x in data.get("deliverables", [])]
            data = clean_scope(db, data)
            scid = new_id()
            payload = json.dumps(data, ensure_ascii=False)
            nr, note = review(sc)
            db.insert("scopes", {"id": scid, "season_id": sid, "data": payload, "doc_version": 1, "needs_review": nr, "review_note": note, "source": "planilha:escopo",
                                 "created_at": ts, "updated_at": ts, "created_by": uid, "updated_by": uid})
            db.insert("scope_versions", {"id": new_id(), "scope_id": scid, "doc_version": 1, "data": payload, "note": "Importado da planilha ESCOPO DE PROJETO", "created_by": uid, "created_at": ts})
            res["scope"]["importados"] += 1
    # evidências
    for w in st["evidence"]:
        if skip("evidence", w):
            continue
        ex = db.one("SELECT id FROM evidence_weeks WHERE season_id=? AND month=? AND week_number=? AND deleted_at IS NULL", (sid, w["month"], w["week_number"]))
        if ex:
            res["evidence"]["duplicados"] += 1
            continue
        db.ex("DELETE FROM evidence_weeks WHERE season_id=? AND month=? AND week_number=? AND deleted_at IS NOT NULL", (sid, w["month"], w["week_number"]))
        wid = new_id()
        nr, note = review(w)
        db.insert("evidence_weeks", {"id": wid, "season_id": sid, "month": w["month"], "week_number": w["week_number"], "needs_review": nr, "review_note": note,
                                     "source": "planilha:" + w["key"], "created_at": ts, "updated_at": ts, "created_by": uid})
        res["evidence"]["importados"] += 1
        res["evidence"]["para_revisar"] += nr
        for e in w["entries"]:
            db.insert("evidence_entries", {"id": new_id(), "week_id": wid, "area_id": area_id(e["area"]), "activities_done": e["text"], "status": "registrado",
                                           "source": "planilha", "created_at": ts, "updated_at": ts, "created_by": uid})
            res["evidence_entries"]["importados"] += 1
    # validação pós-gravação: totais e relações
    checks = list(st.get("checks", []))

    def chk(label, expected, actual):
        checks.append({"check": label, "expected": expected, "actual": actual, "ok": expected == actual})
    exp_acts = sum(1 for a in st["activities"] if a.get("action") == "importar" and not a.get("exists"))
    chk("Atividades gravadas × previstas na revisão", exp_acts, res["activities"]["importados"])
    chk("Elementos da EAP gravados × previstos", sum(1 for w in st["wbs"] if w.get("action") == "importar" and not w.get("exists")), res["wbs"]["importados"])
    orphans = db.val("SELECT COUNT(*) FROM wbs_elements w WHERE w.season_id=? AND w.deleted_at IS NULL AND w.parent_id IS NOT NULL AND NOT EXISTS (SELECT 1 FROM wbs_elements p WHERE p.id=w.parent_id AND p.deleted_at IS NULL)", (sid,))
    chk("Elementos da EAP com elemento superior inexistente", 0, orphans)
    exp_deps = sum(len(a["deps"]) for a in st["activities"] if a["origin"] == "gantt" and a.get("action") == "importar" and not a.get("exists"))
    chk("Dependências gravadas × lidas da planilha", exp_deps, res["dependencies"]["importados"] + res["dependencies"]["duplicados"] + res["dependencies"]["ignorados"])
    chk("Dependências não importadas", 0, res["dependencies"]["ignorados"])
    exp_raci = sum(len([k for k in (a.get("raci") or {}) if k in pmap]) for a in st["activities"] if a.get("action") == "importar" and not a.get("exists"))
    chk("Atribuições RACI gravadas × lidas", exp_raci, res["raci"]["importados"])
    multi_a = db.val("SELECT COUNT(*) FROM (SELECT activity_id FROM raci r JOIN activities a ON a.id=r.activity_id WHERE a.season_id=? AND r.role='A' GROUP BY activity_id HAVING COUNT(*)>1)", (sid,))
    chk("Atividades com mais de um aprovador", 0, multi_a)
    teacher_a = db.val("SELECT COUNT(*) FROM raci r JOIN users u ON u.id=r.user_id JOIN activities a ON a.id=r.activity_id WHERE a.season_id=? AND r.role='A' AND u.is_teacher=1", (sid,))
    chk("Aprovações atribuídas ao professor técnico", 0, teacher_a)
    chk("Movimentos de caixa criados pela importação (nada é lançado como pago/recebido sem comprovação)", 0,
        db.val("SELECT COUNT(*) FROM movements WHERE season_id=? AND created_at >= ?", (sid, ts)))
    return res, checks


# ----------------------------------------------------------------------------- rotas
def _ser_import(i, full=True):
    out = {k: i[k] for k in ("id", "created_at", "status", "season_id", "confirmed_at", "created_by")}
    out["files"] = json.loads(i["files"])
    out["summary"] = json.loads(i["summary"]) if i["summary"] else None
    if full:
        out["staging"] = json.loads(i["staging"])
    return out


def create_import(ctx, season_id, uploads):
    db = ctx.db
    season = db.one("SELECT * FROM seasons WHERE id=? AND deleted_at IS NULL", (season_id,))
    if not season:
        raise bad("Selecione a temporada de destino.")
    if not uploads:
        raise bad("Envie ao menos uma planilha (.xlsx).")
    for up in uploads:
        if not up["name"].lower().endswith(".xlsx") or not up["data"].startswith(b"PK"):
            raise bad(f"“{up['name']}” não é uma planilha .xlsx válida.")
    st = build_staging(db, season, uploads)
    iid = new_id()
    files = []
    for up, info in zip(uploads, st["files"]):
        if up.get("file_id"):   # arquivo já guardado pelo envio em etapas: apenas vincula à importação
            fid = up["file_id"]
            db.ex("UPDATE files SET entity_type='import', entity_id=? WHERE id=?", (iid, fid))
        else:
            fid = new_id()
            db.insert("files", {"id": fid, "entity_type": "import", "entity_id": iid, "name": up["name"][:180], "mime": XLSX_MIME,
                                "size": len(up["data"]), "sha256": info["sha256"], "data": up["data"], "judge_visible": 0, "uploaded_by": ctx.uid, "created_at": now_iso()})
        files.append({**info, "file_id": fid})
    db.insert("imports", {"id": iid, "created_by": ctx.uid, "created_at": now_iso(), "status": "em_revisao", "season_id": season_id,
                          "files": json.dumps(files, ensure_ascii=False), "staging": json.dumps(st, ensure_ascii=False)})
    audit(ctx, "importar_planilhas", "imports", iid, f"{len(uploads)} planilha(s) lida(s) para revisão")
    return _ser_import(db.one("SELECT * FROM imports WHERE id=?", (iid,)))


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@route("POST", "/api/imports/arquivo", allow=("admin",))
def upload_import_file(ctx):
    """Envio em etapas: guarda UMA planilha por requisição (hospedagens limitam o tamanho de cada envio).
    Depois, POST /api/imports com {season_id, file_ids} lê todas juntas."""
    import hashlib
    if len(ctx.files) != 1:
        raise bad("Envie uma planilha por vez.")
    f = ctx.files[0]
    if not f["name"].lower().endswith(".xlsx") or not f["data"].startswith(b"PK"):
        raise bad(f"“{f['name']}” não é uma planilha .xlsx válida.")
    fid = new_id()
    ctx.db.insert("files", {"id": fid, "entity_type": "import_draft", "entity_id": ctx.uid, "name": f["name"][:180], "mime": XLSX_MIME, "size": len(f["data"]),
                            "sha256": hashlib.sha256(f["data"]).hexdigest(), "data": f["data"], "judge_visible": 0, "uploaded_by": ctx.uid, "created_at": now_iso()})
    return {"file_id": fid, "name": f["name"], "size": len(f["data"])}


@route("POST", "/api/imports", allow=("admin",))
def upload_import(ctx):
    ids = ctx.body.get("file_ids")
    if ids:
        if not isinstance(ids, list) or len(ids) > 40 or not all(isinstance(x, str) for x in ids):
            raise bad("Lista de arquivos inválida.")
        ups = []
        for fid in ids:
            row = ctx.db.one("SELECT id, name, data FROM files WHERE id=? AND entity_type='import_draft' AND uploaded_by=? AND deleted_at IS NULL", (fid, ctx.uid))
            if not row:
                raise bad("Um dos arquivos enviados não foi encontrado. Envie as planilhas novamente.")
            ups.append({"name": row["name"], "data": bytes(row["data"]), "file_id": row["id"]})
        return create_import(ctx, ctx.body.get("season_id"), ups)
    return create_import(ctx, ctx.body.get("season_id"), [{"name": f["name"], "data": f["data"]} for f in ctx.files])


@route("GET", "/api/imports", allow=("admin",))
def list_imports(ctx):
    return {"items": [_ser_import(i, False) for i in ctx.db.q("SELECT * FROM imports ORDER BY created_at DESC LIMIT 50")]}


@route("GET", "/api/imports/{id}", allow=("admin",))
def get_import(ctx):
    i = ctx.db.one("SELECT * FROM imports WHERE id=?", (ctx.params["id"],))
    if not i:
        raise not_found("Importação")
    return _ser_import(i)


@route("PUT", "/api/imports/{id}", allow=("admin",))
def update_import(ctx):
    """Ajusta as decisões da revisão: {decisions: {chave_do_item: 'importar'|'ignorar'|'mesclar'}}."""
    i = ctx.db.one("SELECT * FROM imports WHERE id=?", (ctx.params["id"],))
    if not i:
        raise not_found("Importação")
    if i["status"] != "em_revisao":
        raise conflict("Esta importação já foi finalizada.")
    st = json.loads(i["staging"])
    dec = ctx.body.get("decisions") or {}
    for sec in ("people", "wbs", "activities", "risks", "expenses", "stakeholders", "evidence"):
        for it in st[sec]:
            if it["key"] in dec:
                d_ = dec[it["key"]]
                if d_ not in ("importar", "ignorar", "mesclar") or (d_ == "mesclar" and not it.get("merge_into")):
                    raise bad("Decisão inválida para um dos itens.")
                it["action"] = d_
    if st.get("scope") and "escopo" in dec and dec["escopo"] in ("importar", "ignorar"):
        st["scope"]["action"] = dec["escopo"]
    st["summary"] = summarize(st)
    ctx.db.ex("UPDATE imports SET staging=? WHERE id=?", (json.dumps(st, ensure_ascii=False), i["id"]))
    return _ser_import(ctx.db.one("SELECT * FROM imports WHERE id=?", (i["id"],)))


def confirm(ctx, import_id):
    db = ctx.db
    i = db.one("SELECT * FROM imports WHERE id=?", (import_id,))
    if not i:
        raise not_found("Importação")
    if i["status"] != "em_revisao":
        raise conflict("Esta importação já foi finalizada.")
    st = json.loads(i["staging"])
    res, checks = apply_import(ctx, i, st)
    summary = {"result": res, "checks": checks, "labels": {**SECTION_LABEL, "dependencies": "Dependências", "raci": "Atribuições RACI", "evidence_entries": "Registros de evidência"}}
    db.ex("UPDATE imports SET status='confirmada', confirmed_at=?, summary=? WHERE id=?", (now_iso(), json.dumps(summary, ensure_ascii=False), i["id"]))
    total = sum(v["importados"] for v in res.values())
    audit(ctx, "confirmar_importacao", "imports", i["id"], f"Importação confirmada: {total} registro(s) gravado(s)", after={k: v for k, v in res.items()})
    return _ser_import(db.one("SELECT * FROM imports WHERE id=?", (i["id"],)))


@route("POST", "/api/imports/{id}/confirmar", allow=("admin",))
def confirm_import(ctx):
    return confirm(ctx, ctx.params["id"])


@route("POST", "/api/imports/{id}/descartar", allow=("admin",))
def discard_import(ctx):
    i = ctx.db.one("SELECT * FROM imports WHERE id=?", (ctx.params["id"],))
    if not i:
        raise not_found("Importação")
    if i["status"] != "em_revisao":
        raise conflict("Esta importação já foi finalizada.")
    ctx.db.ex("UPDATE imports SET status='descartada' WHERE id=?", (i["id"],))
    audit(ctx, "descartar_importacao", "imports", i["id"], "Importação descartada (os arquivos originais foram mantidos)")
    return {"ok": True}
