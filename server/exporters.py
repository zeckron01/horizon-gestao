"""Renderização de relatórios em PDF (reportlab), Excel (openpyxl) e Word (python-docx) a partir de um modelo único."""
import io
from pathlib import Path

STATIC = Path(__file__).parent / "static"


def _hex(c):
    c = (c or "#147d14").lstrip("#")
    return tuple(int(c[i:i + 2], 16) for i in (0, 2, 4))


def cell_text(v):
    if v is None or v == "":
        return "—"
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, float):
        return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


# ----------------------------------------------------------------------------- PDF
def to_pdf(doc, logo=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    primary = colors.Color(*[x / 255 for x in _hex(doc.get("primary"))])
    wide = doc.get("landscape", False)
    page = landscape(A4) if wide else A4
    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor("#1f2937"))
    small = ParagraphStyle("s", parent=base, fontSize=8, leading=10.5)
    cellh = ParagraphStyle("ch", parent=small, fontName="Helvetica-Bold", textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=20, leading=24, alignment=TA_CENTER, textColor=colors.HexColor("#0f172a"))
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=12, spaceAfter=5, textColor=primary)
    center = ParagraphStyle("c", parent=base, alignment=TA_CENTER, textColor=colors.HexColor("#475569"))
    note = ParagraphStyle("n", parent=small, textColor=colors.HexColor("#64748b"), spaceBefore=3)

    def P(text, st=base):
        return Paragraph(escape(cell_text(text)).replace("\n", "<br/>"), st)

    footer_text = f"{doc.get('team', '')} — {doc['title']} — gerado em {doc['generated_at']} — Fonte: {doc['source']}"

    def on_page(canvas, d_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(15 * mm, 9 * mm, footer_text[:170])
        canvas.drawRightString(page[0] - 15 * mm, 9 * mm, f"Página {canvas.getPageNumber()}")
        canvas.setStrokeColor(primary)
        canvas.setLineWidth(1.2)
        canvas.line(15 * mm, page[1] - 12 * mm, page[0] - 15 * mm, page[1] - 12 * mm)
        canvas.restoreState()

    pdf = SimpleDocTemplate(buf, pagesize=page, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=17 * mm, bottomMargin=16 * mm,
                            title=doc["title"], author=doc.get("team", ""), subject=doc.get("period_label", ""))
    avail = page[0] - 30 * mm
    max_lines = int((page[1] - 33 * mm - 60) / 10.5 * 0.8)  # linhas de texto que cabem com folga em uma página

    def chunks(text, width):
        """Divide textos muito longos para que nenhuma linha de tabela fique mais alta que a página."""
        text = cell_text(text)
        cpl = max(6, int((width - 8) / 4.4))
        out, cur, used = [], [], 0
        for line in text.split("\n"):
            while True:
                need = max(1, -(-len(line) // cpl))
                if used + need <= max_lines:
                    cur.append(line)
                    used += need
                    break
                room = (max_lines - used) * cpl
                if room >= cpl:
                    cut = line.rfind(" ", 0, room)
                    cut = cut if cut > room // 2 else room
                    cur.append(line[:cut])
                    line = line[cut:].lstrip()
                out.append("\n".join(cur))
                cur, used = [], 0
                if not line:
                    break
        if cur or not out:
            out.append("\n".join(cur))
        return out

    def split_rows(rows, widths):
        res = []
        for r in rows:
            parts = [chunks(v, w_) for v, w_ in zip(r, widths)]
            for k in range(max(len(x) for x in parts)):
                res.append([x[k] if k < len(x) else "" for x in parts])
        return res

    flow = []
    # capa
    flow.append(Spacer(1, 26 * mm if not wide else 8 * mm))
    try:
        img = Image(io.BytesIO(logo) if logo else str(STATIC / "logo.png"))
        ratio = img.imageHeight / float(img.imageWidth)
        img.drawWidth = 62 * mm
        img.drawHeight = 62 * mm * ratio
        img.hAlign = "CENTER"
        flow.append(img)
    except Exception:
        pass
    flow += [Spacer(1, 8 * mm), P(doc["title"], h1), Spacer(1, 3 * mm), P(doc.get("team", ""), center)]
    if doc.get("subtitle"):
        flow.append(P(doc["subtitle"], center))
    flow.append(Spacer(1, 8 * mm))
    meta = [["Temporada", doc.get("season", "—")], ["Período de referência", doc.get("period_label", "—")], ["Data de geração", doc["generated_at"]],
            ["Gerado por", doc.get("generated_by", "—")], ["Origem dos dados", doc["source"]]] + [[k, v] for k, v in doc.get("filters", [])]
    t = Table([[P(a, small), P(b, small)] for a, b in meta], colWidths=[45 * mm, 95 * mm], hAlign="CENTER")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    flow += [t, PageBreak()]
    for sec in doc["sections"]:
        flow.append(P(sec["heading"], h2))
        txt_ = sec.get("text")
        for para in ([txt_] if isinstance(txt_, str) else (txt_ or [])):
            flow.append(P(para))
            flow.append(Spacer(1, 2 * mm))
        if sec.get("kv"):
            rows = [[P(k, small), P(v, small)] for k, v in split_rows(sec["kv"], [avail * 0.38, avail * 0.62])]
            t = Table(rows, colWidths=[avail * 0.38, avail * 0.62])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f8fafc")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            flow += [t, Spacer(1, 3 * mm)]
        tb = sec.get("table")
        if tb:
            if not tb["rows"]:
                flow.append(P(tb.get("empty", "Nenhum registro no período e filtros selecionados."), note))
            else:
                n = len(tb["columns"])
                w = tb.get("widths") or [1] * n
                tot = float(sum(w))
                widths = [avail * x / tot for x in w]
                data = [[P(c, cellh) for c in tb["columns"]]] + [[P(v, small) for v in r] for r in split_rows(tb["rows"], widths)]
                t = Table(data, colWidths=widths, repeatRows=1)
                st = [("BACKGROUND", (0, 0), (-1, 0), primary), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                      ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                      ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
                t.setStyle(TableStyle(st))
                flow.append(t)
        if sec.get("note"):
            flow.append(P(sec["note"], note))
        if sec.get("page_break"):
            flow.append(PageBreak())
    pdf.build(flow, onFirstPage=lambda c, d_: None, onLaterPages=on_page)
    return buf.getvalue()


# ----------------------------------------------------------------------------- Excel
def to_xlsx(doc):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    color = (doc.get("primary") or "#147d14").lstrip("#").upper()
    head_fill = PatternFill("solid", fgColor=color)
    ws["A1"] = doc["title"]
    ws["A1"].font = Font(bold=True, size=14)
    meta = [("Equipe", doc.get("team")), ("Temporada", doc.get("season")), ("Período de referência", doc.get("period_label")),
            ("Data de geração", doc["generated_at"]), ("Gerado por", doc.get("generated_by")), ("Origem dos dados", doc["source"])] + list(doc.get("filters", []))
    r = 3
    for k, v in meta:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1
    r += 1
    used = {"Resumo"}
    for sec in doc["sections"]:
        ws.cell(r, 1, sec["heading"]).font = Font(bold=True, size=12, color=color)
        r += 1
        txt_ = sec.get("text")
        for para in ([txt_] if isinstance(txt_, str) else (txt_ or [])):
            c = ws.cell(r, 1, para)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            ws.row_dimensions[r].height = min(300, 15 * (1 + len(str(para)) // 110))
            r += 1
        for k, v in sec.get("kv") or []:
            ws.cell(r, 1, k).font = Font(bold=True)
            ws.cell(r, 2, v if isinstance(v, (int, float)) and not isinstance(v, bool) else cell_text(v))
            r += 1
        tb = sec.get("table")
        if tb:
            name = "".join(ch for ch in sec["heading"] if ch not in '[]:*?/\\')[:28] or "Dados"
            base, i = name, 2
            while name in used:
                name = f"{base[:25]} {i}"
                i += 1
            used.add(name)
            sh = wb.create_sheet(name)
            for ci, col in enumerate(tb["columns"], 1):
                c = sh.cell(1, ci, col)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head_fill
                c.alignment = Alignment(wrap_text=True, vertical="center")
            widths = [len(str(c)) for c in tb["columns"]]
            for ri, row in enumerate(tb["rows"], 2):
                for ci, v in enumerate(row, 1):
                    if isinstance(v, bool):
                        v = "Sim" if v else "Não"
                    c = sh.cell(ri, ci, v if isinstance(v, (int, float)) else (None if v in (None, "") else str(v)))
                    if isinstance(v, float):
                        c.number_format = '#,##0.00'
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    widths[ci - 1] = max(widths[ci - 1], min(60, len(str(v)) if v is not None else 0))
            for ci, wdt in enumerate(widths, 1):
                sh.column_dimensions[get_column_letter(ci)].width = max(10, min(60, wdt + 2))
            sh.freeze_panes = "A2"
            if tb["rows"]:
                sh.auto_filter.ref = sh.dimensions
            ws.cell(r, 1, f"Ver aba “{name}” ({len(tb['rows'])} linha(s)).").font = Font(italic=True, color="64748B")
            r += 1
        if sec.get("note"):
            ws.cell(r, 1, sec["note"]).font = Font(italic=True, color="64748B")
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 48
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------- Word
def to_docx(doc, logo=None):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor
    d = Document()
    sec0 = d.sections[0]
    sec0.page_width, sec0.page_height = Mm(210), Mm(297)
    if doc.get("landscape"):
        sec0.orientation = WD_ORIENT.LANDSCAPE
        sec0.page_width, sec0.page_height = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin"):
        setattr(sec0, side, Mm(18))
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    rgb = RGBColor(*_hex(doc.get("primary")))
    try:
        d.add_picture(io.BytesIO(logo) if logo else str(STATIC / "logo.png"), width=Mm(48))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        pass
    t = d.add_heading(doc["title"], 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(doc.get("team", ""))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = [("Temporada", doc.get("season")), ("Período de referência", doc.get("period_label")), ("Data de geração", doc["generated_at"]),
            ("Gerado por", doc.get("generated_by")), ("Origem dos dados", doc["source"])] + list(doc.get("filters", []))
    tb = d.add_table(rows=0, cols=2)
    tb.style = "Light Grid Accent 1"
    for k, v in meta:
        cells = tb.add_row().cells
        cells[0].text, cells[1].text = str(k), cell_text(v)
    for sec in doc["sections"]:
        h = d.add_heading(sec["heading"], 1)
        for run in h.runs:
            run.font.color.rgb = rgb
        txt_ = sec.get("text")
        for para in ([txt_] if isinstance(txt_, str) else (txt_ or [])):
            d.add_paragraph(str(para))
        if sec.get("kv"):
            tb = d.add_table(rows=0, cols=2)
            tb.style = "Light Grid Accent 1"
            for k, v in sec["kv"]:
                cells = tb.add_row().cells
                cells[0].text, cells[1].text = str(k), cell_text(v)
        t_ = sec.get("table")
        if t_:
            if not t_["rows"]:
                d.add_paragraph(t_.get("empty", "Nenhum registro no período e filtros selecionados.")).italic = True
            else:
                tb = d.add_table(rows=1, cols=len(t_["columns"]))
                tb.style = "Light Grid Accent 1"
                for i, c in enumerate(t_["columns"]):
                    tb.rows[0].cells[i].text = str(c)
                for row in t_["rows"]:
                    cells = tb.add_row().cells
                    for i, v in enumerate(row):
                        cells[i].text = cell_text(v)
                for row in tb.rows:
                    for c in row.cells:
                        for pp in c.paragraphs:
                            for run in pp.runs:
                                run.font.size = Pt(8.5)
        if sec.get("note"):
            n = d.add_paragraph(str(sec["note"]))
            for run in n.runs:
                run.italic = True
                run.font.size = Pt(8.5)
    foot = sec0.footer.paragraphs[0]
    foot.text = f"{doc.get('team', '')} — {doc['title']} — gerado em {doc['generated_at']} — Fonte: {doc['source']}"
    for run in foot.runs:
        run.font.size = Pt(7.5)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
